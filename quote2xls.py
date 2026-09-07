#!/usr/bin/env python3
"""
quote2xls.py — turn supplier quotation PDFs into structured Excel/CSV.

Extracts line items (part number, qty, UOM, description, unit price, line
total) from text-based quote PDFs, cross-checks the arithmetic, and flags
anomalies (totals that do not multiply, missing prices, tax/freight lines).

No API keys. No network. Works offline. Best results on table-based quotes
(bordered or cleanly aligned columns). Scanned/image PDFs are out of scope
(no OCR) and are reported as such.

Usage:
    python quote2xls.py quote.pdf [more.pdf ...] [-o outdir] [--csv]

Output per input file (in outdir/, default ./out):
    <name>.xlsx   sheet "Items" (metadata + flagged rows), sheet "Report"
    <name>.csv    line items only (UTF-8 BOM, opens clean in Excel)
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    import pymupdf  # modern name (1.24+); PyMuPDF
except ImportError:
    import fitz as pymupdf  # older alias

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required:  pip install openpyxl")

# --------------------------------------------------------------------------
# number parsing
# --------------------------------------------------------------------------

def parse_number(raw: str) -> float | None:
    """Parse '1,234.56', '12.50', '1.234,56', 'CAD 45.00', '-'. None if not numeric."""
    if raw is None:
        return None
    s = str(raw).strip()
    if not s or s in ("-", "--", "—", "N/A", "n/a", ""):
        return None
    s = re.sub(r"[^\d.,+\-]", "", s)
    if not s or s in ("-", "+"):
        return None
    # European style: 1.234,56  -> decimal comma with thousand dots
    if "," in s and "." in s and s.rfind(",") > s.rfind("."):
        s = s.replace(".", "").replace(",", ".")
    elif "," in s and "." not in s and len(s.split(",")[-1]) == 2 and len(s.split(",")[0]) > 3:
        s = s.replace(",", ".")
    else:
        s = s.replace(",", "")
    try:
        f = float(s)
        return f if f >= 0 else None   # negative totals/qty are never legit on a quote
    except ValueError:
        return None


# --------------------------------------------------------------------------
# document metadata (best effort)
# --------------------------------------------------------------------------

HEADER_KW = {"qty", "quantity", "unit", "price", "total", "amount", "part",
             "item", "description", "desc", "make", "model", "uom", "no."}
META_KW = {
    "qty": ("qty", "quantity", "qty."),
    "uom": ("uom", "unit", "units"),
    "part": ("part", "part no", "item", "no.", "code"),
    "desc": ("description", "desc", "item description", "product", "make", "model"),
    "price": ("price", "unit price", "rate", "cost"),
    "total": ("total", "amount", "line total", "extended"),
}
SUMMARY_RE = re.compile(
    r"^(sub\s*total|total|balance|amount\s*(due|paid|owing)|gst|hst|qst|pst|"
    r"vat|tax|freight|shipping|handling|discount|deposit|env|levy)\b",
    re.IGNORECASE,
)

MONTH_RE = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*"


def detect_supplier(text: str) -> str:
    lines = [l.strip() for l in text.splitlines() if l.strip()][:14]
    strong = re.compile(r"(?i)(\b(?:inc|ltd|llc|corp|corporation|co\.|limited|"
                        r"gmbh|s\.a\.|pty)\b|@[\w.]+\.\w{2,4})")
    for ln in lines:
        if strong.search(ln) and len(ln) < 90 and not re.search(r"(?i)quote|quotation|invoice", ln):
            return ln
    return ""


def detect_quote_no(text: str) -> str:
    # line-scoped: avoid matching a bare "QUOTATION" heading
    for line in text.splitlines():
        if not re.search(r"(?i)\b(?:quote|quotation|rfq|offer|reference|po)\b", line):
            continue
        m = re.search(r"(?i)\b(?:quote|quotation|rfq|offer|reference|po)\b", line)
        tail = line[m.end():]
        tail = re.sub(r"(?i)^[\s:#.\-]*(?:no\.?|number)?[\s:#.\-]*", "", tail)
        tok = re.match(r"([A-Z0-9][A-Z0-9\-_/. ]{2,24}?)(?:\s|$)", tail)
        if tok and re.search(r"\d", tok.group(1)):
            return tok.group(1).strip()
    return ""


def detect_date(text: str) -> str:
    m = re.search(r"\b(20\d{2})[-/.](\d{1,2})[-/.](\d{1,2})\b", text)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    m = re.search(rf"\b{MONTH_RE}\s+\d{{1,2}},?\s+(20\d{{2}})\b", text, re.IGNORECASE)
    if m:
        return m.group(0)
    m = re.search(rf"\b(\d{{1,2}})[-/.]({MONTH_RE})[-/.](\d{{4}})\b", text, re.IGNORECASE)
    if m:
        return f"{m.group(1)}-{m.group(2).capitalize()}-{m.group(3)}"
    m = re.search(r"\b(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})\b", text)
    if m:
        return f"{int(m.group(3)):04d}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    return ""


def detect_currency(text: str) -> str:
    if re.search(r"\bCAD\b|\bC\$|CA\$|\$", text) and not re.search(r"\bUS\$|\bUSD\b", text):
        return "CAD"
    if re.search(r"\bUSD\b|\bUS\$", text):
        return "USD"
    if "€" in text or re.search(r"\bEUR\b", text):
        return "EUR"
    if "£" in text or re.search(r"\bGBP\b", text):
        return "GBP"
    return ""


# --------------------------------------------------------------------------
# table parsing
# --------------------------------------------------------------------------

def map_columns(header: list[str]) -> dict[str, int]:
    """Map header cells to roles. Returns {role: col_index}."""
    out: dict[str, int] = {}
    used: set[int] = set()
    for role, kws in META_KW.items():
        for i, cell in enumerate(header):
            if i in used:
                continue
            c = re.sub(r"[^a-z]", "", str(cell).lower())
            if c in kws or any(k in c for k in kws):
                out[role] = i
                used.add(i)
                break
    return out


def _row_has_data(row: list[str], cols: dict[str, int]) -> bool:
    for role in ("qty", "price", "total"):
        i = cols.get(role)
        if i is not None and i < len(row) and parse_number(row[i]) is not None:
            return True
    return False


def assemble(part: str, qty_raw, uom: str, desc: str, price_raw, total_raw) -> dict:
    """Build one item dict + math/anomaly flags from raw cell strings."""
    qty = parse_number(qty_raw) if qty_raw else None
    price = parse_number(price_raw) if price_raw else None
    total = parse_number(total_raw) if total_raw else None
    if qty is None and uom:
        m = re.match(r"([\d.,]+)\s*([A-Za-z]+)", uom)
        if m:
            qty, uom = parse_number(m.group(1)), m.group(2)
    flags: list[str] = []
    if qty is not None and price is not None and total is not None:
        expect = qty * price
        if expect and abs(expect - total) > 0.02 * max(abs(expect), abs(total), 1.0):
            flags.append(f"math: {qty:g} x {price:g} != {total:g}")
    elif total is not None and price is None:
        flags.append("no unit price")
    elif price is not None and qty is None and total is None:
        flags.append("no quantity")
    return {"part": part.strip() if part else "", "qty": qty, "uom": uom,
            "desc": desc.strip() if desc else "", "price": price,
            "total": total, "flags": flags}


def classify_table(tbl_rows: list[list[str]], supplier: str) -> dict:
    """Given extracted rows, find the header and build items + flags."""
    header_idx, header = None, None
    for idx, row in enumerate(tbl_rows[:6]):
        cells = [str(c).strip() for c in row]
        score = sum(1 for c in cells if re.sub(r"[^a-z]", "", c.lower()) in HEADER_KW)
        if score >= 2:
            header_idx, header = idx, cells
            break
    if header is None:
        # headerless table: maybe a standalone totals block (Subtotal/Freight/HST/Total)
        totals: list[dict] = []
        for row in tbl_rows:
            cells = [str(c).strip() for c in row]
            text = " ".join(cells).strip()
            if not text:
                continue
            if SUMMARY_RE.match(text):
                amt = parse_number(cells[-1]) if cells else None
                if amt is not None:
                    totals.append({"text": text[:60], "amount": amt})
        if totals:
            return {"items": [], "totals": totals, "note": ""}
        return {"items": [], "flags": [], "totals": [], "note": "no header row found"}
    cols = map_columns(header)
    items: list[dict] = []
    totals: list[dict] = []
    last = None
    for row in tbl_rows[header_idx + 1:]:
        cells = [str(c).strip() for c in row]
        if not any(cells):
            continue
        text = " ".join(cells).strip()
        if not text or (supplier and supplier.lower() in text.lower()):
            continue
        di = cols.get("desc", 0)
        di = di if di < len(cells) else 0
        first = cells[0] or ""
        check = cells[di] if di < len(cells) else first
        if SUMMARY_RE.match(re.sub(r"^[\d.\s]+$", "", check) or check or ""):
            amt = parse_number(cells[cols["total"]]) if cols.get("total", 99) < len(cells) else None
            totals.append({"text": text[:60], "amount": amt})
            continue
        if not _row_has_data(cells, cols):
            if last and di < len(cells) and cells[di]:
                last["desc"] = (last["desc"] + " " + cells[di]).strip()
            continue
        gi = cols.get("desc", 0)
        desc = cells[gi] if gi < len(cells) else ""
        it = assemble(
            cells[cols["part"]] if cols.get("part", 99) < len(cells) else "",
            cells[cols["qty"]] if cols.get("qty", 99) < len(cells) else "",
            cells[cols["uom"]] if cols.get("uom", 99) < len(cells) else "",
            desc,
            cells[cols["price"]] if cols.get("price", 99) < len(cells) else "",
            cells[cols["total"]] if cols.get("total", 99) < len(cells) else "",
        )
        if it["price"] is None and it["total"] is None and it["qty"] is None and not it["desc"]:
            continue
        items.append(it)
        last = it
    return {"items": items, "totals": totals, "note": ""}


# --------------------------------------------------------------------------
# borderless (aligned, unruled) column parser
# --------------------------------------------------------------------------

def _cluster_words(words) -> list[list]:
    rows: list[list] = []
    cur: list = []
    cur_y: float | None = None
    for w in sorted(words, key=lambda w: (round(w[1] / 2.5), w[0])):
        if cur_y is None or abs(w[1] - cur_y) <= 2.5:
            cur.append(w)
            if cur_y is None:
                cur_y = w[1]
        else:
            rows.append(cur)
            cur = [w]
            cur_y = w[1]
    if cur:
        rows.append(cur)
    return rows


def _role_of_token(tok: str) -> str | None:
    t = re.sub(r"[^a-z]", "", tok.lower())
    if not t:
        return None
    for role, kws in META_KW.items():
        if t in kws:
            return role
    return None


def parse_borderless(page, supplier: str) -> dict:
    """Best-effort column parse for cleanly aligned, unruled tables."""
    words = page.get_text("words")
    if not words:
        return {"items": [], "totals": [], "note": "no words on page"}
    rows = _cluster_words(words)
    hdr_idx = None
    for i, r in enumerate(rows[:12]):
        toks = [w[4] for w in r]
        if sum(1 for t in toks if _role_of_token(t) in ("qty", "uom", "desc", "price", "total", "part")) >= 3:
            hdr_idx = i
            break
    if hdr_idx is None:
        return {"items": [], "totals": [], "note": "no aligned header row"}
    hdr = rows[hdr_idx]
    xs = [w[0] for w in hdr]
    bounds = sorted(set(round(x, 1) for x in xs))
    # column boundaries between distinct header token starts
    boundaries = []
    for i, x in enumerate(bounds):
        if i == 0 or x - bounds[i - 1] > 6:
            boundaries.append(x)
    if len(boundaries) < 2:
        return {"items": [], "totals": [], "note": "columns too narrow"}
    # assign each header token's role to its column start
    col_role = {}
    for w in hdr:
        x0 = w[0]
        col = min(range(len(boundaries)), key=lambda j: abs(boundaries[j] - x0))
        role = _role_of_token(w[4])
        if role and col not in col_role:
            col_role[col] = role
    stops = boundaries + [max(w[2] for w in words) + 40]

    def cell_for(x_center: float) -> int:
        for j in range(len(boundaries)):
            if x_center < boundaries[j]:
                return j - 1
        return len(boundaries) - 1

    items: list[dict] = []
    totals: list[dict] = []
    last = None
    saw_summary = False
    for r in rows[hdr_idx + 1:]:
        cells = {j: [] for j in range(len(boundaries))}
        for w in r:
            cells[cell_for((w[0] + w[2]) / 2)].append(w[4])
        joined = {j: " ".join(v).strip() for j, v in cells.items()}
        full = " ".join(joined.values()).strip()
        if not full or (supplier and supplier.lower() in full.lower()):
            continue
        num_cols = {j: parse_number(v) for j, v in joined.items() if v}
        has_role_num = any(parse_number(v) is not None for j, v in joined.items()
                           if j in col_role and col_role[j] in ("qty", "price", "total"))
        # summary/total rows: text label (+ trailing amount), few tokens
        if len(r) <= 6 and any(v is not None for v in num_cols.values()):
            words_txt = [w[4] for w in r]
            while words_txt and parse_number(words_txt[-1]) is not None:
                words_txt.pop()
            if SUMMARY_RE.match(" ".join(words_txt).strip()):
                amt = next((v for v in num_cols.values() if v is not None), None)
                totals.append({"text": full[:60], "amount": amt})
                saw_summary = True
                continue
        if saw_summary:
            continue  # anything below the totals block is footer, not items
        if not has_role_num:
            # continuation line of the previous item's description
            if last and any(w[4].isalpha() for w in r):
                extras = " ".join(w[4] for w in r).strip()
                last["desc"] = (last["desc"] + " " + extras).strip()
            continue
        role_at = {role: next((j for j, r_ in col_role.items() if r_ == role), None) for role in ("part", "qty", "uom", "desc", "price", "total")}
        overflow = " ".join(w[4] for w in r
                            if (w[0] + w[2]) / 2 < boundaries[0] and parse_number(w[4]) is None)
        it = assemble(
            joined[role_at["part"]] if role_at["part"] is not None else "",
            joined[role_at["qty"]] if role_at["qty"] is not None else "",
            joined[role_at["uom"]] if role_at["uom"] is not None else "",
            (overflow + " " + joined[role_at["desc"]]) if role_at["desc"] is not None and overflow
            else (joined[role_at["desc"]] if role_at["desc"] is not None else ""),
            joined[role_at["price"]] if role_at["price"] is not None else "",
            joined[role_at["total"]] if role_at["total"] is not None else "",
        )
        if it["price"] is None and it["total"] is None and it["qty"] is None and not it["desc"]:
            continue
        items.append(it)
        last = it
    return {"items": items, "totals": totals, "note": ""}


def extract_document(path: Path) -> dict:
    doc = pymupdf.open(path)
    all_text = "\n".join(page.get_text("text") for page in doc)
    meta = {
        "source": path.name,
        "supplier": detect_supplier(all_text),
        "quote_no": detect_quote_no(all_text),
        "date": detect_date(all_text),
        "currency": detect_currency(all_text),
        "pages": len(doc),
    }
    items: list[dict] = []
    totals: list[dict] = []
    notes: list[str] = []
    tables_found = 0
    for page in doc:
        for strategy in ("lines", "text"):
            try:
                found = page.find_tables(strategy=strategy)
            except Exception:
                found = None
            if not found or not found.tables:
                continue
            tables_found += 1
            for t in found.tables[:6]:
                rows = [[str(c).strip() for c in r] for r in t.extract()]
                if len(rows) < 2:
                    continue
                res = classify_table(rows, meta["supplier"])
                if res["note"] and res["note"] not in notes:
                    notes.append(res["note"])
                items.extend(res["items"])
                totals.extend(res["totals"])
    doc.close()
    if not items:
        # fallback: cleanly aligned but unruled layouts
        doc = pymupdf.open(path)
        for page in doc:
            res = parse_borderless(page, meta["supplier"])
            items.extend(res["items"])
            totals.extend(res["totals"])
            if res["note"] and res["note"] not in notes:
                notes.append(f"page {page.number + 1}: {res['note']}")
        doc.close()
        if not items:
            notes.append("no tabular structure detected - scanned/image PDF or free-form layout")
        else:
            # fallback succeeded: drop misleading structured-pass notes
            notes = [n for n in notes if "no header row" not in n]
    # overlapping detections across strategies/tables/passes: keep unique rows
    nz = lambda s: re.sub(r"\s+", "", s or "")
    seen_i, uniq_items = set(), []
    for it in items:
        sig = (it["part"], it["qty"], it["uom"], round(it["price"], 2) if it["price"] is not None else None,
               round(it["total"], 2) if it["total"] is not None else None, nz(it["desc"])[:60])
        if sig in seen_i:
            continue
        seen_i.add(sig)
        uniq_items.append(it)
    seen_t, uniq_tot = set(), []
    for t in totals:
        sig = (round(t["amount"], 2) if t["amount"] is not None else None, nz(t["text"])[:40])
        if sig in seen_t:
            continue
        seen_t.add(sig)
        uniq_tot.append(t)
    return {**meta, "items": uniq_items, "totals": uniq_tot,
            "notes": notes, "tables_found": tables_found}


# --------------------------------------------------------------------------
# output writers
# --------------------------------------------------------------------------

def write_xlsx(result: dict, out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Items"
    ws.append(["Source", result["source"]])
    ws.append(["Supplier", result["supplier"]])
    ws.append(["Quote / Ref", result["quote_no"]])
    ws.append(["Date", result["date"]])
    ws.append(["Currency", result["currency"]])
    ws.append(["Pages", result["pages"]])
    ws.append(["Line items", len(result["items"])])
    if result["notes"]:
        ws.append(["Notes", "; ".join(result["notes"])])
    ws.append([])
    head = ["#", "Part No", "Qty", "UOM", "Description", "Unit Price", "Line Total", "Flags"]
    ws.append(head)
    bold = Font(bold=True)
    for c in range(1, len(head) + 1):
        ws.cell(row=ws.max_row, column=c).font = bold
    flag_fill = PatternFill("solid", fgColor="FFE3E3")
    for i, it in enumerate(result["items"], 1):
        r = ws.max_row + 1
        ws.append([i, it["part"], it["qty"], it["uom"], it["desc"],
                   it["price"], it["total"], "; ".join(it["flags"])])
        if it["flags"]:
            for c in range(1, len(head) + 1):
                ws.cell(row=r, column=c).fill = flag_fill
    ws.append([])
    ws.append(["SUMMARY / tax lines"])
    for c in range(1, len(head) + 1):
        ws.cell(row=ws.max_row, column=c).font = bold
    for t in result["totals"]:
        ws.append(["", "", "", "", t["text"], "", t["amount"], ""])
    widths = [4, 16, 8, 8, 52, 12, 12, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    rep = wb.create_sheet("Report")
    rep.append(["Metric", "Value"])
    rep.append(["Source", result["source"]])
    rep.append(["Supplier detected", result["supplier"]])
    rep.append(["Quote / Ref detected", result["quote_no"]])
    rep.append(["Date detected", result["date"]])
    rep.append(["Currency detected", result["currency"]])
    rep.append(["Tables found", result["tables_found"]])
    rep.append(["Items extracted", len(result["items"])])
    flagged = sum(1 for it in result["items"] if it["flags"])
    rep.append(["Items flagged", flagged])
    rep.append(["Notes", "; ".join(result["notes"]) or "none"])
    rep.column_dimensions["A"].width = 22
    rep.column_dimensions["B"].width = 80
    wb.save(out_path)


def write_csv(result: dict, out_path: Path) -> None:
    with open(out_path, "w", encoding="utf-8-sig", newline="") as f:
        f.write("Part No,Qty,UOM,Description,Unit Price,Line Total,Flags\n")
        for it in result["items"]:
            esc = lambda s: '"' + str(s).replace('"', '""') + '"'
            f.write(",".join([
                esc(it["part"]), str(it["qty"] or ""), esc(it["uom"]),
                esc(it["desc"]), str(it["price"] or ""), str(it["total"] or ""),
                esc("; ".join(it["flags"])),
            ]) + "\n")


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------

def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pdfs", nargs="+", help="quote PDF file(s)")
    ap.add_argument("-o", "--outdir", default="out", help="output directory (default: ./out)")
    ap.add_argument("--csv", action="store_true", help="also write .csv per file")
    a = ap.parse_args(argv)

    outdir = Path(a.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    rc = 0
    for p in a.pdfs:
        src = Path(p)
        if not src.exists():
            print(f"  !! {p}: file not found")
            rc = 1
            continue
        try:
            res = extract_document(src)
        except Exception as e:  # noqa: BLE001
            print(f"  !! {src.name}: extraction failed: {e}")
            rc = 1
            continue
        stem = src.stem
        write_xlsx(res, outdir / f"{stem}.xlsx")
        if a.csv:
            write_csv(res, outdir / f"{stem}.csv")
        flagged = sum(1 for it in res["items"] if it["flags"])
        print(f"  {src.name}: {len(res['items'])} items, {flagged} flagged"
              f" | {res['supplier'] or 'supplier?'} | {res['quote_no'] or 'no ref'}"
              f" | {res['date'] or 'no date'}")
        for it in res["items"]:
            if it["flags"]:
                print(f"      ! row {it['desc'][:48] or it['part']}: {'; '.join(it['flags'])}")
        for n in res["notes"]:
            print(f"      note: {n}")
    print(f"  -> wrote to {outdir.resolve()}")
    return rc


if __name__ == "__main__":
    raise SystemExit(main())
